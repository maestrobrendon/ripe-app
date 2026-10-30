"""
Exports the recipe library sheet to prisma/data/library.json for `npm run library:import`.

    python prisma/scripts/export-library.py [path/to/sheet.xlsx]

Default sheet: docs/Basket_Produce_and_Recipe_Library_v2.xlsx. Needs openpyxl.

The sheet must have been saved by Excel (or recalculated) so its formula cells
carry values: the nutrition, tagline and coverage columns are formulas.

- A recipe already in the current library.json (same sheet ID) keeps its slug,
  goal and sort order, so the import recognises it and re-exporting is stable.
- New recipes get a slug from their title and a goal from their goal tags.
- Ingredients: Basket rows become catalogue slugs (PRODUCT_SLUG below),
  Protein and Pantry rows go to "elsewhere", Staple rows (salt, oil, water)
  are left out, as before.
- Products: the Produce tab did not change in v2, so the products the
  catalogue lacked (already imported in Oct 2026) are carried forward as they are.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SHEET = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "Basket_Produce_and_Recipe_Library_v2.xlsx"
OUT = ROOT / "prisma" / "data" / "library.json"

# Sheet produce ID -> catalogue slug.
PRODUCT_SLUG = {
    "P001": "banana", "P002": "plantain", "P003": "pineapple", "P004": "watermelon", "P005": "pawpaw",
    "P006": "mango", "P007": "orange", "P008": "tangerine", "P009": "lemon", "P010": "lime",
    "P011": "grapefruit", "P012": "avocado", "P013": "guava", "P014": "soursop", "P015": "passion-fruit",
    "P016": "coconut", "P017": "agbalumo", "P018": "cashew-apple", "P019": "velvet-tamarind", "P020": "hog-plum",
    "P021": "pepper-fruit", "P022": "monkey-kola", "P023": "african-pear", "P024": "tiger-nut", "P025": "sugarcane",
    "P026": "custard-apple", "P027": "jackfruit", "P028": "june-plum", "P029": "sweet-melon", "P030": "dates",
    "P031": "apple", "P032": "pear", "P033": "red-grapes", "P034": "strawberries", "P035": "blueberries",
    "P036": "kiwi", "P037": "pomegranate", "P038": "dragon-fruit", "P039": "star-fruit", "P040": "groundnut",
    "P041": "african-walnut", "P042": "tomato", "P043": "cherry-tomato", "P044": "bell-pepper", "P045": "green-bell-pepper",
    "P046": "chilli", "P047": "long-red-chilli", "P048": "onion", "P049": "spring-onion", "P050": "shallots",
    "P051": "garlic", "P052": "ginger", "P053": "turmeric", "P054": "cucumber", "P055": "carrot",
    "P056": "cabbage", "P057": "purple-cabbage", "P058": "romaine-lettuce", "P059": "green-beans", "P060": "corn",
    "P061": "garden-egg", "P062": "okra", "P063": "broccoli", "P064": "cauliflower", "P065": "beetroot",
    "P066": "zucchini", "P067": "mushrooms", "P068": "pumpkin", "P069": "sweet-potato", "P070": "irish-potato",
    "P071": "yam", "P072": "ugu", "P073": "spinach", "P074": "waterleaf", "P075": "baby-spinach",
    "P076": "kale", "P077": "scent-leaf", "P078": "bitter-leaf", "P079": "mint", "P080": "coriander",
    "P081": "parsley", "P082": "basil", "P083": "lemongrass", "P084": "rosemary",
}

# Titles whose natural slug already belongs to a menu recipe.
SLUG_OVERRIDE = {
    "chicken-shawarma-bowl": "chicken-shawarma-salad-bowl",
    # The dotless ı in "Salatası" has no ASCII form, so it would be dropped.
    "shepherds-salad-coban-salatas": "shepherds-salad-coban-salatasi",
}

# Sheet units -> the singular units the app pluralises itself ("pc" means a count).
UNIT = {
    "pc": "", "pcs": "", "": "", "cloves": "clove", "cups": "cup", "stalks": "stalk", "slices": "slice",
    "thumbs": "thumb", "cobs": "cob", "leaves": "leaf", "sprigs": "sprig", "heads": "head",
}

NUTRIENTS = [  # Recipes header, JSON key, % daily label
    ("Fibre (g)", "fibre_g", "fibre"),
    ("Vitamin C (mg)", "vitaminC_mg", "vitaminC"),
    ("Vitamin A (mcg RAE)", "vitaminA_mcgRAE", "vitaminA"),
    ("Vitamin K (mcg)", "vitaminK_mcg", "vitaminK"),
    ("Folate (mcg)", "folate_mcg", "folate"),
    ("Potassium (mg)", "potassium_mg", "potassium"),
    ("Vitamin E (mg)", "vitaminE_mg", "vitaminE"),
    ("Vitamin B6 (mg)", "vitaminB6_mg", "vitaminB6"),
    ("Magnesium (mg)", "magnesium_mg", "magnesium"),
    ("Iron (mg)", "iron_mg", "iron"),
]
PCT_HEADER = {"fibre": "Fibre", "vitaminC": "Vitamin C", "vitaminA": "Vitamin A", "vitaminK": "Vitamin K", "folate": "Folate",
              "potassium": "Potassium", "vitaminE": "Vitamin E", "vitaminB6": "Vitamin B6", "magnesium": "Magnesium", "iron": "Iron"}


def slugify(title: str) -> str:
    s = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    s = s.lower().replace("&", " and ").replace("'", "")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return SLUG_OVERRIDE.get(s, s)


def goal_for(row) -> str:
    """Goal for a recipe new in v2. v1 goals were set by hand and are kept."""
    if row["Category"] == "Fruit Plates & Desserts":
        return "fruit"
    first = (row["Goal tags"] or "").split(",")[0].strip()
    if first in ("Training", "Energy"):
        return "training"
    if first in ("Weight management", "Hydration"):
        return "lighter"
    if first == "Plant-based" or (row["Diet"] == "Vegan" and row["Category"] in ("Stews & Sauces", "Salads & Cooked Sides")):
        return "plants"
    if (row["Total (min)"] or 99) <= 20:
        return "busy"
    return "family"


def num(v, places=1):
    return None if v is None else round(float(v), places)


def main():
    previous = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else {"products": [], "recipes": []}
    v1 = {r["sheetId"]: r for r in previous["recipes"]}

    wb = openpyxl.load_workbook(SHEET, data_only=True)
    ws = wb["Recipes"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    if rows and rows[0]["Basket coverage"] is None:
        raise SystemExit("The sheet's formulas have no saved values. Open it in Excel and save, then run this again.")

    ingredients = {}
    wi = wb["Recipe Ingredients"]
    ih = [c.value for c in wi[1]]
    for r in wi.iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[0]).startswith("R"):
            ingredients.setdefault(r[0], []).append(dict(zip(ih, r)))

    next_sort = max([r["sortOrder"] for r in previous["recipes"]] + [990]) + 10
    recipes, slugs, problems = [], set(), []
    for row in rows:
        rid = row["ID"]
        old = v1.get(rid)
        slug = old["slug"] if old else slugify(row["Recipe"])
        if slug in slugs:
            raise SystemExit(f"Two recipes would share the slug {slug}")
        slugs.add(slug)

        basket, elsewhere, amounts = [], [], {}
        for ing in ingredients.get(rid, []):
            source = ing["Source"]
            if source == "Staple":
                continue
            if source == "Basket":
                pid = ing["Produce ID"]
                if pid not in PRODUCT_SLUG:
                    problems.append(f"{rid} {ing['Ingredient']}: no catalogue slug for {pid}")
                    continue
                key = PRODUCT_SLUG[pid]
                if key not in basket:
                    basket.append(key)
            else:
                key = ing["Ingredient"]
                if key not in elsewhere:
                    elsewhere.append(key)
            qty = ing["Qty"]
            if isinstance(qty, (int, float)):
                unit = UNIT.get((ing["Unit"] or "").strip(), (ing["Unit"] or "").strip())
                if key in amounts and amounts[key][1] == unit:
                    amounts[key][0] = round(amounts[key][0] + float(qty), 2)
                elif key not in amounts:
                    amounts[key] = [float(qty), unit]

        steps = [re.sub(r"^\s*\d+[.)]\s*", "", s).strip() for s in (row["Method"] or "").split("\n")]
        nutrition = {
            "perServing": {key: num(row[h]) for h, key, _ in NUTRIENTS},
            "pctDaily": {label: num(row[f"{PCT_HEADER[label]} % daily"], 2) for _, _, label in NUTRIENTS},
            "strongest": row["Strongest nutrient"] or None,
            "produceGrams": round(row["Produce (g) / serving"] or 0),
            "kcal": row["Est. kcal / serving"],
            "protein_g": row["Est. protein (g) / serving"],
        }
        recipes.append({
            "sheetId": rid,
            "slug": slug,
            "title": row["Recipe"],
            "summary": row["Local adaptation"] or "",
            "goal": old["goal"] if old else goal_for(row),
            "minutes": int(row["Total (min)"] or 0),
            "serves": int(row["Servings"] or 1),
            "basketSlugs": basket,
            "elsewhere": elsewhere,
            "amounts": amounts,
            "steps": [s for s in steps if s],
            "tip": row["Tip"] or None,
            "sortOrder": old["sortOrder"] if old else next_sort,
            "category": row["Category"],
            "subtype": row["Subtype"] or None,
            "serveWith": row["Serve with / you add"] or None,
            "sourceSite": row["Source"] or None,
            "sourceUrl": row["Source URL"] or None,
            "referenceImageUrl": row["Reference image URL"] or None,
            "nutrition": nutrition,
            "healthTagline": row["Health tagline"] or None,
            "alsoStrongIn": row["Also strong in"] or None,
            "healthNote": row["Health note"] or None,
        })
        if not old:
            next_sort += 10

    if problems:
        raise SystemExit("\n".join(problems))

    dropped = [{"sheetId": r["sheetId"], "slug": r["slug"], "title": r["title"]} for r in previous["recipes"] if r["sheetId"] not in {x["sheetId"] for x in recipes}]
    dropped += previous.get("dropped", [])
    seen = set()
    dropped = [d for d in dropped if not (d["slug"] in seen or seen.add(d["slug"]))]

    out = {
        "sheet": SHEET.name,
        "products": previous["products"],
        "recipes": recipes,
        # In the v1 library, out of v2 (drinks, soups, plantain-led plates). The import leaves them as drafts.
        "dropped": dropped,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    new = sum(1 for r in recipes if r["sheetId"] not in v1)
    print(f"{OUT.relative_to(ROOT)}: {len(recipes)} recipes ({len(recipes) - new} already in library.json, {new} new), "
          f"{len(out['products'])} products, {len(dropped)} earlier recipes no longer in the library.")


if __name__ == "__main__":
    main()
